"""Extract plain text from PDF / DOCX / TXT uploads."""
import io

from docx import Document
from pypdf import PdfReader


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Use PDF, DOCX or TXT.")
    text = text.strip()
    if not text:
        raise ValueError("No readable text found (scanned PDFs need OCR).")
    return text
